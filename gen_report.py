# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# --- Stiller ---
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)

NAVY = RGBColor(0x1F, 0x30, 0x4D)
RED = RGBColor(0xD3, 0x20, 0x27)

def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY
    return h

def code(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(9)
    return p

def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p

def table_from(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9.5)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    return t

# ============ BAŞLIK ============
title = doc.add_heading('KalmanFilterComparison — Detaylı İnceleme Raporu', level=0)
for run in title.runs:
    run.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = sub.add_run('GNSS/IMU füzyonu ile Kalman filtre varyantlarının karşılaştırması — '
                'tren (Modena/RTK) ve telefon (Marmaray/Sensor Logger) veri setleri')
r.italic = True
r.font.size = Pt(10.5)

meta = doc.add_paragraph()
meta.add_run('Repo: ').bold = True
meta.add_run('github.com/RSMlabGroup22/KalmanFilterComparison (upstream)\n')
meta.add_run('İncelenen sürüm: ').bold = True
meta.add_run('commit 696fa40 (2026-05-09)\n')
meta.add_run('Rapor tarihi: ').bold = True
meta.add_run('2026-05-30')

doc.add_paragraph()

# ============ 1 ============
heading('1. Genel Mimari ve Veri Akışı', 1)
code(
"data/  veya  data_marmaray/\n"
"        |\n"
"        v\n"
"loadAndSyncSensors.m  -->  ds (time, dt, ax/ay/az, gx/gy/gz,\n"
"                               gps.{lat,lon,h,hacc,heading,time})\n"
"        |\n"
"        +--> KF.m            (Lineer KF)\n"
"        +--> SageHusaKF.m    (Adaptif KF)\n"
"        +--> STF.m           (BOZUK -> imuLoc icerigi)\n"
"        +--> EKF.m           (Genisletilmis KF)\n"
"        +--> EKF_sagehusa.m  (Adaptif EKF)\n"
"        +--> imuLoc.m        (Saf INS / Dead-Reckoning)\n"
"        |\n"
"        v\n"
"plotResults*.m --> Plot/*.png  +  *_predictions.csv\n"
"main.m: calistirir -> cizdirir -> CSV yazar -> 20 tekrarla SURE olcer"
)
doc.add_paragraph('Aktif olmayan (orphan) dosyalar: EKF_strong_tracking.m, ned2lla_local.m '
                  '(main.m içinden çağrılmıyor).')

# ============ 2 ============
heading('2. Veri Katmanı — loadAndSyncSensors.m', 1)
doc.add_paragraph('İki veri formatı otomatik algılanıyor:')
table_from(
    ['', 'Eski (Modena / tren)', 'Yeni (Marmaray / telefon)'],
    [
        ['Tetik', 'clean_imu.csv + clean_gnss.csv', 'Accelerometer.csv varsa'],
        ['Kaynak', 'u-blox RTK GNSS + endüstriyel IMU (~1776 Hz)', 'Telefon — Sensor Logger uygulaması'],
        ['İvme', "g'li (az ≈ 9.8)", 'lineer (yerçekimsiz), ×9.81 ile ölçekleniyor'],
        ['Zaman', 't - T0', 'seconds_elapsed'],
    ],
    widths=[0.8, 2.6, 2.8]
)
doc.add_paragraph('Sorunlar:')
bullet(' ds.dt = mean(diff(time)) → sabit dt varsayımı; telefon örneklemesi jitter\'lı.')
bullet(' speed, velN/velE, altitude sütunları Location.csv\'de var ama okunmuyor.')
bullet(' Loader yorumu ("Z\'ye 9.81 ekliyoruz") kodla uyuşmuyor (kod sadece ×9.81 yapıyor).')
bullet(' Marmaray horizontalAccuracy ilk satırlarda ≈ 1400 m — telefon GPS\'i başlangıçta '
       'çok kötü; outlier reddi olmadığı için filtreyi bozar.')

# ============ 3 ============
heading('3. Algoritma Modülleri', 1)

heading('3.1 KF — KF.m  [sağlam]', 2)
bullet('[N, E, vN, vE]; ivme kontrol girişi (B·u), yaw dışarıda gz entegrasyonuyla.', 'Durum: ')
bullet('sadece konum (lat/lon→NED); R dinamik (hacc²).', 'Ölçüm: ')
bullet('temiz, doğru lineer KF.', 'Artı: ')
bullet('hız hiç ölçülmüyor (konumdan dolaylı), bias yok, yaw filtre dışı.', 'Eksi: ')

heading('3.2 SageHusaKF — SageHusaKF.m', 2)
bullet('KF + Q/R\'nin çevrimiçi adaptasyonu (unutma faktörü b=0.98).')
bullet('R güncellemesi negatif tanımlı olabilir; max(...,1e-4) ile zorlanıyor (kırılgan).', 'Eksi: ')

heading('3.3 STF — STF.m  [BOZUK — EN KRİTİK BULGU]', 2)
p = doc.add_paragraph()
p.add_run('Dosya adı STF ama içeriği function imu_out = imuLoc(ds). Strong Tracking kodu '
          'silinmiş, yerine bir dead-reckoning kopyası yazılmış.').font.color.rgb = RED
bullet('MATLAB dosyayı dosya adıyla çağırdığı için STF(dataStruct) bu DR kodunu çalıştırır.')
bullet('Kod ds.lat(1) / ds.lon(2) kullanıyor; ama loader ds.gps.lat üretiyor — ds.lat alanı yok '
       '→ main.m STF adımında çalışma-zamanı hatası verir ("Unrecognized field \'lat\'").')
bullet('Sonuç: Strong Tracking Filter fiilen kayıp; plotResultsSTF ve stf_predictions.csv anlamsız. '
       'Doğru STF, backup/local-asb-f41d2dc branch\'inde duruyor.')

heading('3.4 EKF — EKF.m', 2)
bullet('[N, E, v, yaw]; ileri-hız modeli, F Jacobian\'ı doğru.', 'Durum: ')
pp = doc.add_paragraph()
pp.add_run('Telefon verisiyle uyumsuz: satır 20 a_res = sqrt(ax²+ay²+(az-9.81)²). Marmaray verisi '
           'yerçekimsiz (az≈0) olduğundan (az-9.81)≈-9.81 → a_res sürekli ~9.81 şişer → EKF '
           'marmarayda hatalı (Modena\'da doğru).').font.color.rgb = RED

heading('3.5 EKF_sagehusa — EKF_sagehusa.m', 2)
bullet('EKF + R adaptasyonu. a_res = ds.ax(k) (yerçekimi çıkarmıyor) → EKF.m\'deki gravity '
       'hatasından etkilenmiyor, ama EKF ile tutarsız (farklı ivme tanımı).')

heading('3.6 EKF_strong_tracking — EKF_strong_tracking.m  [olgun ama orphan]', 2)
bullet('Aslında en olgun filtre: chi-square outlier kapısı (gate_limit=11.3), hard GPS reddi '
       '(hacc>15), ZUPT benzeri durdurma, Joseph-form kovaryans, açı normalizasyonu.')
bullet('Ama main.m\'de hiç çağrılmıyor. Telefon GPS\'i için en gerekli özellikler burada '
       'duruyor ve kullanılmıyor.')

heading('3.7 imuLoc — imuLoc.m  [asıl hedef: ivmeden hız/konum]', 2)
bullet('Saf INS / dead-reckoning: ivme→hız→konum, GPS\'siz (yalnız başlangıç GPS\'ten).')
bullet('5 Hz Butterworth low-pass (inline) + yaw entegrasyonu + v_max=38.5 m/s hız kelepçesi.')
bullet('bias kestirimi yok, ZUPT yok, çalışırken GPS düzeltmesi yok → konum sınırsız sürüklenir; '
       'v_max kelepçesi gerçek drift kontrolü değil (38.5 m/s tren hızı, yaya için anlamsız).', 'Eksi: ')
bullet('İki farklı imuLoc var: bu dosya (Butterworth+clamp, ds.gps.lat) ile STF.m içine düşen '
       'kopya (Euler 0.5at², clamp yok, ds.lat) — birbiriyle çelişiyor.', 'Dikkat: ')

# ============ 4 ============
heading('4. main + Görselleştirme — main.m', 1)
bullet('5 filtre + imuLoc çalıştırır, CSV ve PNG üretir.')
bullet('STF adımında muhtemelen çöker (bkz. 3.3).')
bullet('Deney döngüsü yalnızca çalışma süresini ölçüyor; doğruluk metriği (RMSE/CEP) yok — '
       'referansa kıyas yapılmıyor. "Comparison" projesi için temel eksik.')
bullet('plotResultsimu.m çıktıyı pos_comparison.png\'e yazıyor → başka plot\'larla dosya adı '
       'çakışması (üzerine yazma) riski.')

# ============ 5 ============
heading('5. Kritik Bulgular (öncelik sırasıyla)', 1)
table_from(
    ['#', 'Önem', 'Bulgu', 'Dosya'],
    [
        ['1', 'KRİTİK', 'STF.m bozuk; STF kayıp + runtime hatası (ds.lat)', 'STF.m'],
        ['2', 'KRİTİK', 'EKF yerçekimsiz telefon verisinde (az-9.81) yüzünden hatalı', 'EKF.m:20'],
        ['3', 'YÜKSEK', 'En iyi filtre (outlier-gate\'li) main\'de çağrılmıyor', 'EKF_strong_tracking.m'],
        ['4', 'YÜKSEK', 'Doğruluk metriği (RMSE) yok; sadece süre ölçülüyor', 'main.m'],
        ['5', 'YÜKSEK', 'GPS hızı (Doppler) okunmuyor/füzyona katılmıyor', 'loadAndSync + filtreler'],
        ['6', 'ORTA', 'İki çelişkili imuLoc; bias/ZUPT yok, drift sınırsız', 'imuLoc.m, STF.m'],
        ['7', 'ORTA', 'Sabit dt; ivme ×9.81 birim belirsizliği', 'loadAndSync'],
        ['8', 'DÜŞÜK', 'ned2lla kod tekrarı; 2B\'ye sınırlı, altitude kullanılmıyor', 'tüm filtreler'],
    ],
    widths=[0.4, 0.9, 3.6, 1.8]
)

# ============ 6 ============
heading('6. Telefon Hedefine Uygunluk', 1)
doc.add_paragraph('Hedef: basit telefon GPS + ivmeden ayrı ayrı hız ve konum çıkarmak.')
bullet('Telefon formatı (Sensor Logger) okunuyor.', 'VAR: ')
bullet('Saf-ivme DR modülü (imuLoc) mevcut.', 'VAR: ')
bullet('Drift kontrolü yok (bias/ZUPT); GPS hızı kullanılmıyor; EKF bozuk; STF çalışmıyor; '
       'sayısal doğruluk ölçülmüyor; telefon eğimi (roll/pitch) modellenmiyor.', 'YOK: ')
p = doc.add_paragraph()
p.add_run('Net durum: ').bold = True
p.add_run('Proje "telefon" yönünde doğru başlamış ama şu an çalışır bütünlükte değil '
          '(STF crash, EKF telefonda hatalı) ve hedef için kritik parçalar (drift\'i sınırlayan '
          'füzyon + metrik) eksik.')

# ============ 7 ============
heading('7. Önerilen Yol Haritası', 1)
steps = [
    'STF.m\'i onar — doğru Strong Tracking kodunu backup/local-asb-f41d2dc\'ten geri getir (~5 dk).',
    'EKF gravity\'sini formata göre koşullu yap (Modena\'da çıkar, telefonda çıkarma).',
    'EKF_strong_tracking\'i main\'e dahil et (outlier-gate telefon GPS\'i için şart).',
    'GPS Doppler hızını Location.csv\'den oku, H\'ye hız ölçümü ekle → hız gözlemlenebilirliği + drift sınırlama.',
    'RMSE/CEP metriği ekle (GPS referans alıp tüm filtreleri sayısal kıyasla).',
    'imuLoc\'u tekilleştir + bias durumu / ZUPT ekle.',
]
for i, s in enumerate(steps, 1):
    p = doc.add_paragraph(style='List Number')
    p.add_run(s)

doc.add_paragraph()
note = doc.add_paragraph()
note.add_run('Hızlı kazanç: ').bold = True
note.add_run('İlk iki madde (STF onarımı + EKF gravity) projeyi "çalışır" hale getirir.')

out = r'c:\Users\RSMLAB1\Documents\GitHub\srhtbynkln\KalmanFilterComparison\KalmanFilter_Inceleme_Raporu.docx'
doc.save(out)
print('OK:', out)
